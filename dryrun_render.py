"""Does --render lay out, fit and check an application, and --publish ship it, with no Word?

No Word, no network, no account. Word is replaced by a stand-in that reads the
.docx the renderer wrote and lays its paragraphs out as a real PDF - fixed line
height, wrapped at a fixed width, paragraph spacing honoured - so the fitting,
the line counting and every check run on genuine PDF text. The web is a dict of
URLs, and the portfolio repo a throwaway git repo with a bare "remote".

  * the starter spec reads a master resume in the shape the real ones use,
  * the fit keeps the fullest page that still fits, and reports how far over a
    page that does not,
  * a bullet that runs to three lines is named,
  * the checks catch British spelling, em dashes, a previous employer's name, a
    letter that never names this employer, a broken link, a long subject, and a
    career.yaml rule - and a clean application passes with none,
  * each board writes exactly its files, and never a .docx,
  * --publish pushes the PDF under a random name, logs it, waits for the live
    copy to match, and fills the link into the message.

    python dryrun_render.py
"""

import os
import shutil
import subprocess
import tempfile
import textwrap
from pathlib import Path

import docx
import yaml

from jobsift import render as R
from jobsift.publish import publish

# See dryrun_evidence.py: a monitor daemon per throwaway repo once hung the suite.
os.environ.setdefault("GIT_CONFIG_COUNT", "1")
os.environ.setdefault("GIT_CONFIG_KEY_0", "core.fsmonitor")
os.environ.setdefault("GIT_CONFIG_VALUE_0", "false")
for key, value in (("GIT_AUTHOR_NAME", "dryrun"), ("GIT_AUTHOR_EMAIL", "dryrun@example.com"),
                   ("GIT_COMMITTER_NAME", "dryrun"), ("GIT_COMMITTER_EMAIL", "dryrun@example.com")):
    os.environ.setdefault(key, value)

failures = 0


def check(label, ok, detail=""):
    global failures
    if not ok:
        failures += 1
    print(f"  {'ok  ' if ok else 'FAIL'}  {label:<56} {str(detail)[:110]}")


# ── A stand-in for Word ──────────────────────────────────────────────────────


def tiny_pdf(pages):
    """A real PDF, one Tj per line at the given heights."""
    def esc(text):
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    objs = {1: b"<< /Type /Catalog /Pages 2 0 R >>",
            3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"}
    kids, n = [], 4
    for lines in pages:
        stream = ("BT /F1 10 Tf " + " ".join(f"1 0 0 1 50 {y:.2f} Tm ({esc(t)}) Tj"
                                             for y, t in lines) + " ET").encode("latin-1")
        objs[n] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        objs[n + 1] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources "
                       f"<< /Font << /F1 3 0 R >> >> /Contents {n} 0 R >>").encode()
        kids.append(n + 1)
        n += 2
    objs[2] = (f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] "
               f"/Count {len(kids)} >>").encode()
    out, offsets = bytearray(b"%PDF-1.4\n"), {}
    for i in sorted(objs):
        offsets[i] = len(out)
        out += f"{i} 0 obj\n".encode() + objs[i] + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for i in range(1, len(objs) + 1):
        out += f"{offsets[i]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


BOTTOM = R.MARGINS_IN["bottom"] * 72
TOP = 792 - R.MARGINS_IN["top"] * 72 - 12


def fake_word(pairs):
    for src, dst in pairs:
        pages, lines, y = [], [], TOP
        for p in docx.Document(str(src)).paragraphs:
            if not p.text.strip():
                continue
            fmt = p.paragraph_format
            y -= fmt.space_before.pt if fmt.space_before else 0
            for chunk in textwrap.wrap(p.text, 95) or [p.text]:
                if y < BOTTOM:
                    pages.append(lines)
                    lines, y = [], TOP
                lines.append((y, chunk))
                y -= 12
            y -= fmt.space_after.pt if fmt.space_after else 0
        pages.append(lines)
        dst.write_bytes(tiny_pdf(pages))


WEB = {}


class Response:
    def __init__(self, status, content=b""):
        self.status_code, self.content = status, content


def fake_fetch(url):
    return WEB.get(url, Response(200))


# ── Fixtures ─────────────────────────────────────────────────────────────────

PROFILE = {"name": "Ana Cruz", "email": "ana@example.com", "phone": "+63 900 000 0000",
           "location": "Manila", "portfolio": "https://ana.example.dev",
           "github": "https://github.com/anacruz"}
RULES = [{"id": "players-with-rate", "when": ["570 players"], "require": ["on a given day"],
          "why": "a lifetime total travels with its current rate"}]
LETTER = """Dear Hiring Team,

I am applying for the Backend Developer role at Initech. I build the part of a product that has to keep working when nobody is watching it.

At my current contract I am the only engineer on a billing platform fifteen people use every day. I wrote the schema, the access rules and the admin screens, and I answer for all of it. When a nurse could read a colleague's pay rate, I found it, closed it and wrote down how it happened.

Your posting asks for Postgres and background jobs. Both are daily work for me. What I have not done is lead a team, and I would rather say that plainly than stretch a title to cover it.

I can start immediately, and I work remotely already with a team in Manila and an owner in the United States. My work, with the code behind it, is on my portfolio.

Sincerely,
Ana Cruz"""


def bullets(n, words=14):
    return [f"Built thing number {i} with a clear constraint and an outcome someone can check "
            + "and more words " * max(0, words - 14) for i in range(n)]


def spec(company, board="none", sections=None, **extra):
    data = {"company": company, "company_names": [], "employer": f"{company} role", "board": board,
            "letter": "", "resume": {"headline": "Backend Developer  |  Postgres  |  AI",
                                     "sections": sections or [
                                         {"heading": "Summary", "bullets": bullets(3)},
                                         {"heading": "Experience", "entries": [
                                             {"title": "Acme - Engineer", "detail": "Jun 2026 - Present",
                                              "bullets": bullets(6)}]}]}}
    data.update(extra)
    return data


def write_spec(folder, data):
    path = Path(folder) / f"{data['company']}.yaml"
    path.write_text(yaml.safe_dump(data, sort_keys=False, allow_unicode=True), encoding="utf-8")
    return path


work = Path(tempfile.mkdtemp(prefix="jobsift-render-dry-"))
try:
    specs, out = work / "specs", work / "out"
    specs.mkdir()
    out.mkdir()
    SETTINGS = R.settings_for({"output_dir": str(out), "specs_dir": str(specs)})

    def run(data, **kw):
        return R.render(write_spec(specs, data), PROFILE, SETTINGS, RULES, convert=fake_word,
                        fetch=fake_fetch, **kw)

    print("the spec\n" + "-" * 78)
    try:
        R.load_spec(write_spec(specs, {"company": "Two Words", "board": "email", "resume": {}}))
        check("a bad spec is refused", False)
    except R.SpecError as exc:
        text = str(exc)
        check("a bad spec is refused, every problem at once",
              all(k in text for k in ("one token", "headline", "sections", "letter text",
                                      "email.subject")), text.replace("\n", " "))

    master = work / "master.docx"
    doc = docx.Document()
    doc.add_paragraph().add_run("ANA CRUZ").bold = True
    doc.add_paragraph("Backend Developer | Builder\nManila | ana@example.com")
    doc.add_paragraph().add_run("\nPROFESSIONAL SUMMARY").bold = True
    doc.add_paragraph("Five years of shipped work.", style="List Bullet")
    doc.add_paragraph().add_run("EXPERIENCE").bold = True
    doc.add_paragraph().add_run("Acme - Engineer (Contract) | June 2026 - Present").bold = True
    doc.add_paragraph("Sole engineer on the billing platform.", style="List Bullet")
    doc.add_paragraph("Built the admin screens.", style="List Bullet")
    doc.add_paragraph().add_run("PROJECTS").bold = True
    doc.add_paragraph().add_run("Jobsift - Job Alert Pipeline").bold = True
    doc.add_paragraph("Reads five sources.", style="List Bullet")
    doc.save(str(master))
    started = yaml.safe_load(R.starter_spec(master, "Initech"))
    secs = started["resume"]["sections"]
    check("the starter keeps the headline", started["resume"]["headline"] == "Backend Developer | Builder",
          started["resume"]["headline"])
    check("and the section headings", [s["heading"] for s in secs] == ["PROFESSIONAL SUMMARY", "EXPERIENCE", "PROJECTS"],
          [s["heading"] for s in secs])
    check("an entry's right-hand part goes to `detail`",
          secs[1]["entries"][0] == {"title": "Acme - Engineer (Contract)", "detail": "June 2026 - Present",
                                    "bullets": ["Sole engineer on the billing platform.", "Built the admin screens."]},
          secs[1]["entries"][0])
    check("a project without a date keeps its whole title",
          secs[2]["entries"][0]["title"] == "Jobsift - Job Alert Pipeline")
    check("section bullets stay on the section", secs[0]["bullets"] == ["Five years of shipped work."])
    # The AI master's shape: title-case headings, the detail in a run of its own.
    # The first starter read none of its sections, because it wanted capitals.
    titled = work / "titled.docx"
    doc = docx.Document()
    doc.add_paragraph().add_run("ANA CRUZ").bold = True
    doc.add_paragraph("AI Developer  |  Automation")
    doc.add_paragraph("Manila  |  ana@example.com")
    doc.add_paragraph().add_run("Professional Summary").bold = True
    doc.add_paragraph("Five years.", style="List Bullet")
    doc.add_paragraph().add_run("Projects").bold = True
    p = doc.add_paragraph()
    p.add_run("RecodeAI - AI Website Redesigner").bold = True
    p.add_run("   |   Claude API")
    doc.add_paragraph("Crawls a site.", style="List Bullet")
    doc.save(str(titled))
    t = yaml.safe_load(R.starter_spec(titled, "Initech"))["resume"]
    check("a title-case master reads the same way",
          [s["heading"] for s in t["sections"]] == ["Professional Summary", "Projects"]
          and t["sections"][1]["entries"][0] == {"title": "RecodeAI - AI Website Redesigner",
                                                 "detail": "Claude API", "bullets": ["Crawls a site."]},
          t["sections"])
    starter_path = specs / "Initech.yaml"
    starter_path.write_text(R.starter_spec(master, "Initech"), encoding="utf-8")
    check("the starter is a spec that loads as written", R.load_spec(starter_path)["company"] == "Initech")

    print("\nthe fit\n" + "-" * 78)
    short = run(spec("Short"))
    check("a short resume uses the loosest spacing", short["stretch"] == max(R.STRETCHES), short["stretch"])
    check("and says how much room is left", any("room for about" in w for w in short["warn"]), short["warn"])
    # Grow the resume a job at a time until the loosest spacing no longer fits:
    # that one must still fit one page, at the fullest spacing that does.
    medium = None
    for jobs in range(2, 12):
        got = run(spec("Medium", sections=[{"heading": "Summary", "bullets": bullets(3)},
                                           {"heading": "Experience", "entries": [
                                               {"title": f"Job {j}", "detail": "2020", "bullets": bullets(6)}
                                               for j in range(jobs)]}]))
        if got["stretch"] < max(R.STRETCHES):
            medium = got
            break
    check("a nearly full resume is tightened only as far as it must be",
          medium is not None and medium["pages"] == 1 and medium["stretch"] > min(R.STRETCHES),
          medium and (medium["pages"], medium["stretch"], medium["waste_in"]))
    check("and then the page is full", medium is not None and medium["waste_in"] <= R.MAX_WASTE_IN,
          medium and medium["waste_in"])
    long = run(spec("Long", sections=[{"heading": "Experience", "entries": [
        {"title": f"Job {j}", "bullets": bullets(8)} for j in range(8)]}]))
    check("too much for a page fails, with how far over",
          long["pages"] > 1 and any("line(s) over" in f for f in long["fail"]), long["fail"][:1])
    wrapped = run(spec("Wrapped", sections=[{"heading": "Summary",
                                             "bullets": bullets(2) + bullets(1, words=45)}]))
    check("a bullet that runs to three lines is named",
          any("runs to" in w and "keep it to two" in w for w in wrapped["warn"]), wrapped["warn"])
    check("a two-line bullet is not",
          sum("runs to" in w for w in wrapped["warn"]) == 1)
    check("no .docx is ever left in the output folder", not list(out.glob("*.docx")),
          [p.name for p in out.glob("*.docx")])

    print("\nthe checks\n" + "-" * 78)
    write_spec(specs, spec("Globex", company_names=["Globex Corp"], letter="Dear Globex Corp, hi."))
    (out / "Ana_Cruz_Resume_SparkMarketing.pdf").write_bytes(b"old")
    WEB["https://ana.example.dev/broken"] = Response(404)
    clean = run(spec("Initech", board="email", company_names=["Initech"], letter=LETTER,
                     email={"to": "jobs@initech.example", "subject": "Backend Developer - Ana Cruz"}))
    check("a clean application has nothing to fix", clean["fail"] == [], clean["fail"])
    email = (out / "Ana_Cruz_Email_Initech.txt").read_text(encoding="utf-8")
    check("the email starts with its address and subject",
          email.startswith("TO: jobs@initech.example\nSUBJECT: Backend Developer - Ana Cruz\n\n"))
    check("the sign-off is profile.yaml's, not the model's",
          "Yours truly,\nAna Cruz" in email and "Sincerely" not in email, email[-80:])
    check("email board: the resume PDF and the email, nothing else",
          sorted(p.name for p in out.glob("Ana_Cruz_*Initech*"))
          == ["Ana_Cruz_Email_Initech.txt", "Ana_Cruz_Resume_Initech.pdf"])

    dirty_sections = [{"heading": "Summary", "bullets": [
        "Optimised the billing sync and 570 players joined a game.",
        "Rebuilt the portfolio at ana.example.dev/broken for Spark Marketing."]}]
    dirty = run(spec("Umbrella", board="onlinejobs", company_names=["Umbrella"], sections=dirty_sections,
                     letter=LETTER.replace("Initech", "Globex Corp").replace("plainly", "plainly — honestly"),
                     message={"subject": "S" * 90}))
    fails = " | ".join(dirty["fail"])
    check("British spelling fails", "British spelling: Optimised" in fails, fails[:120])
    check("an em dash fails", "em dash" in fails)
    check("an earlier employer's name fails", "'Globex Corp', an employer from an earlier" in fails)
    check("a letter that never names this employer fails", "never names the employer (Umbrella)" in fails)
    check("a broken link fails", "ana.example.dev/broken: HTTP 404" in fails)
    check("a long subject fails", "message.subject: 90 characters" in fails)
    check("a career.yaml rule fails", "players-with-rate" in fails)
    warns = " | ".join(dirty["warn"])
    check("a name from the output folder is flagged, not failed",
          "'Spark Marketing', which is also an earlier" in warns, warns[:120])
    check("the message waits for its link", "no resume link yet" in warns)
    message = (out / "Ana_Cruz_Message_Umbrella.txt").read_text(encoding="utf-8")
    check("the message ends on the link line, no sign-off",
          message.rstrip().endswith(R.LINK_PLACEHOLDER) and "Yours truly" not in message)
    allowed = run(spec("Allowed", sections=[{"heading": "Summary", "bullets": ["Built the Optimised Centre site."]}],
                       allow_words=["Optimised", "Centre"]))
    check("allow_words lets a proper noun through", not any("British" in f for f in allowed["fail"]),
          allowed["fail"])
    offline = run(spec("Offline", sections=dirty_sections), offline=True)
    check("--offline checks no links", not any("HTTP" in f for f in offline["fail"])
          and not offline["links_checked"])
    form = run(spec("Formco", board="form", company_names=["Formco"], letter=LETTER.replace("Initech", "Formco")))
    check("form board: a cover letter PDF beside the resume",
          sorted(p.name for p in out.glob("Ana_Cruz_*Formco*"))
          == ["Ana_Cruz_CoverLetter_Formco.pdf", "Ana_Cruz_Resume_Formco.pdf"])

    print("\n--publish\n" + "-" * 78)
    origin, repo = work / "origin.git", work / "port"
    git = lambda *a, cwd=None: subprocess.run(["git", *a], cwd=cwd, capture_output=True, text=True, check=True)
    git("init", "-q", "--bare", str(origin))
    git("init", "-q", str(repo))
    (repo / "docs").mkdir()
    (repo / "docs" / "resume-links.md").write_text(
        "# Links\n\n| Sent | Employer | URL | Built from |\n|---|---|---|---|\n", encoding="utf-8")
    git("add", "-A", cwd=repo)
    git("commit", "-q", "-m", "start", cwd=repo)
    git("remote", "add", "origin", str(origin), cwd=repo)
    git("push", "-q", "-u", "origin", "HEAD", cwd=repo)
    (repo / "notes.txt").write_text("unrelated work in progress\n", encoding="utf-8")

    job = spec("Umbrella", board="onlinejobs", company_names=["Umbrella"],
               letter=LETTER.replace("Initech", "Umbrella"), message={"subject": "Backend Developer"},
               employer="Backend Developer, Umbrella (onlinejobs.ph)")
    spec_file = write_spec(specs, job)
    run(job)
    pdf = out / "Ana_Cruz_Resume_Umbrella.pdf"
    settings = {**SETTINGS, "publish_repo": str(repo), "publish_url": "https://ana.example.dev/r/"}
    live = lambda url: Response(200, pdf.read_bytes()) if "/r/umbrella-" in url else Response(404)
    done = publish(spec_file, PROFILE, settings, fetch=live, today="2026-09-11", poll_every=0)
    name = done["url"].rsplit("/", 1)[-1]
    check("published under a random name", done["url"].startswith("https://ana.example.dev/r/umbrella-")
          and len(name) == len("umbrella-") + 8 + len(".pdf"), done["url"])
    check("the file is in r/ and byte-identical", (repo / "r" / name).read_bytes() == pdf.read_bytes())
    log = (repo / "docs" / "resume-links.md").read_text(encoding="utf-8")
    check("logged with date, employer and URL",
          f"| 2026-09-11 | Backend Developer, Umbrella (onlinejobs.ph) | `/r/{name}` |" in log)
    pushed = subprocess.run(["git", "--git-dir", str(origin), "log", "-1", "--format=%s", "--name-only"],
                            capture_output=True, text=True).stdout
    check("pushed as its own commit, with only its two files",
          "Publish the Umbrella resume variant" in pushed and f"r/{name}" in pushed
          and "notes.txt" not in pushed, pushed.replace("\n", " "))
    check("the live copy was checked", done["verified"] is True)
    message = (out / "Ana_Cruz_Message_Umbrella.txt").read_text(encoding="utf-8")
    check("the link is filled into the message", message.rstrip().endswith(f"Resume: {done['url']}"))
    check("and into the spec, for a re-render",
          (yaml.safe_load(spec_file.read_text(encoding="utf-8"))["message"] or {}).get("link") == done["url"])
    rerun = run(yaml.safe_load(spec_file.read_text(encoding="utf-8")))
    check("a re-render keeps the link", not any("no resume link" in w for w in rerun["warn"])
          and (out / "Ana_Cruz_Message_Umbrella.txt").read_text(encoding="utf-8").rstrip().endswith(done["url"]))
    stale = publish(spec_file, PROFILE, settings, fetch=lambda url: Response(200, b"old deploy"),
                    wait_seconds=0, poll_every=0)
    check("a live file that does not match is reported", stale["verified"] is False)
finally:
    shutil.rmtree(work, ignore_errors=True)

print("\n" + "-" * 78 + f"\n{'all cases behaved' if not failures else f'{failures} FAILED'}")
raise SystemExit(1 if failures else 0)

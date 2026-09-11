"""Render a tailored application from a content file: code lays it out, the model writes it.

The job-application-tailoring skill used to have the agent write a python-docx
builder from scratch for every posting - about 100 lines, 50 of them the same
page setup each time - then convert it in Word, measure the page, edit and go
round again, three to six times per application. On 2026-09-10 that loop was
about 55% of an 89-request OpenCode run. Everything in it except the words is
deterministic, so it lives here:

  * the layout and type scale (10.5pt body, one column, no tables), and the
    header from profile.yaml, so a phone number is never retyped;
  * the fit: paragraph spacing is tried across a range in ONE Word session and
    the fullest one-page result kept, so nobody tunes spacing by trial;
  * every check an application has to pass, in one report.

The model writes a spec (`--render-init COMPANY` starts one from the master
resume) and reads the report. What is left to it is what needs judgment: which
facts, which words, and what to cut when the report says a line over.
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import tempfile
import time
import unicodedata
from pathlib import Path

import yaml

DEFAULTS = {
    "output_dir": "./applications",       # where finished PDFs and texts go
    "specs_dir": "./data/applications",   # one spec per application, kept
    "masters": [],                        # resume .docx files --render-init starts from
    "publish_repo": "",                   # a git repo that serves r/ (see publish.py)
    "publish_url": "",                    # the URL r/ is served at
    "publish_log": "docs/resume-links.md",
}

# How each board receives an application, which decides what gets written.
#   email       resume PDF + the email (TO, SUBJECT, letter, sign-off) as text
#   onlinejobs  resume PDF + subject and message ending on the resume link, no
#               sign-off: the board appends the name and has a contact box
#   form        resume PDF + a cover letter PDF when there is a letter
#   none        resume PDF only
BOARDS = ("email", "onlinejobs", "form", "none")

# The type scale, from the variants shipped on 2026-09-10 (see the skill, §0a).
FONT = "Calibri"
SIZES = {"name": 20, "headline": 11.5, "contact": 9.5, "heading": 11.5,
         "entry": 10.5, "entry_detail": 10, "body": 10.5, "letter": 11}
MARGINS_IN = {"top": 0.42, "bottom": 0.42, "left": 0.55, "right": 0.55}
# Paragraph spacing in points at stretch 1.0. Only these move when fitting; the
# type never shrinks, because a full page at a readable size beats a cramped one.
SPACING = {"header_after": 2, "heading_before": 4.5, "heading_after": 2,
           "entry_before": 2.5, "entry_after": 1, "bullet_after": 1.5}
STRETCHES = (0.55, 0.7, 0.85, 1.0, 1.15, 1.3, 1.45, 1.6, 1.8)
# A resume page is full when less than this is left under the last line.
MAX_WASTE_IN = 0.25
LETTER_WORDS = (150, 220)
SUBJECT_MAX = 80
LINK_PLACEHOLDER = "Resume: <fill in with --publish>"

# British forms, as whole words. Not the skill's old grep, which matched
# "realistic" on "realis": these need the ending that makes them British.
BRITISH = re.compile(
    r"\b(?:(?:summar|organ|optim|docker|priorit|recogn|real|util|standard|minim|maxim|"
    r"custom|normal|author|capital|categor|emphas|special|visual|final|stabil|monet)"
    r"is(?:e|es|ed|ing|ation|ations|er|ers)"
    r"|analys(?:e|es|ed|ing)"
    r"|(?:behavi|col|lab|fav|hon|hum|neighb|flav)our(?:s|ed|ing|ite|ites)?"
    r"|licence|defence|offence|whilst|judgement|programme|programmes"
    r"|centre|centres|datacentre|catalogue"
    r"|(?:model|travel|cancel|label|signal|fuel)l(?:ed|ing))\b",
    re.I,
)


class SpecError(Exception):
    """A spec that cannot be rendered, with every problem found, not just the first."""


class RenderError(Exception):
    """Something outside the spec failed: no way to make a PDF, a missing file."""


def settings_for(config_render: dict | None) -> dict:
    return {**DEFAULTS, **(config_render or {})}


def file_prefix(profile: dict) -> str:
    return "_".join(str(profile.get("name") or "Resume").split())


def spec_path(arg: str, settings: dict) -> Path:
    """`--render` takes a spec file or a company token."""
    path = Path(arg)
    if path.suffix in (".yaml", ".yml") or path.exists():
        return path
    return Path(settings["specs_dir"]) / f"{arg}.yaml"


# ── The spec ─────────────────────────────────────────────────────────────────


def load_spec(path) -> dict:
    try:
        spec = yaml.safe_load(Path(path).read_text(encoding="utf-8")) or {}
    except FileNotFoundError:
        raise SpecError(f"No spec at {path}. Start one with --render-init COMPANY.")
    problems = []
    company = str(spec.get("company") or "")
    if not re.fullmatch(r"[A-Za-z0-9]+", company):
        problems.append("company: one token, letters and digits (WhiteCloak, not White Cloak)")
    board = spec.get("board") or "none"
    if board not in BOARDS:
        problems.append(f"board: one of {', '.join(BOARDS)}")
    resume = spec.get("resume") or {}
    if not str(resume.get("headline") or "").strip():
        problems.append("resume.headline is empty")
    sections = resume.get("sections") or []
    if not sections:
        problems.append("resume.sections is empty")
    for n, section in enumerate(sections, 1):
        if not isinstance(section, dict) or not str(section.get("heading") or "").strip():
            problems.append(f"resume.sections[{n}] needs a heading")
            continue
        if not (section.get("bullets") or section.get("entries")):
            problems.append(f"section {section['heading']!r} has no bullets and no entries")
        for entry in section.get("entries") or []:
            if not isinstance(entry, dict) or not str(entry.get("title") or "").strip():
                problems.append(f"an entry under {section['heading']!r} has no title")
    letter = str(spec.get("letter") or "").strip()
    if board in ("email", "onlinejobs") and not letter:
        problems.append(f"board {board} needs the letter text under `letter:`")
    if board == "email" and not str((spec.get("email") or {}).get("subject") or "").strip():
        problems.append("board email needs email.subject")
    if board == "onlinejobs" and not str((spec.get("message") or {}).get("subject") or "").strip():
        problems.append("board onlinejobs needs message.subject")
    if letter and not spec.get("company_names"):
        problems.append("company_names: how the letter names the employer, e.g. [Asset Opportunity AI]")
    if problems:
        raise SpecError("The spec is not ready:\n" + "\n".join(f"  - {p}" for p in problems))
    spec["board"] = board
    return spec


SPEC_HELP = """\
# One application. The words are yours; layout, fit and checks are jobsift's.
#   python -m jobsift --render {company}
#
# company        one token, names the files: Kim_Julongbayan_Resume_{company}.pdf
# company_names  how the letter names the employer; the letter must use one, and
#                no later application may (the stale-name check)
# employer       one line for the published-links log, e.g. "AI Developer, Acme (email)"
# board          email | onlinejobs | form | none - decides what else is written:
#                  email       the email as text: email.to, email.subject, letter
#                  onlinejobs  message.subject + letter, ending on the resume link
#                  form        a cover letter PDF, when there is a letter
#                  none        the resume only
# letter         greeting and paragraphs, blank line between them. No sign-off:
#                it is added from profile.yaml where the board needs one.
# allow_words    words the spelling check must let through (a proper noun: Centre)
# resume         headline, then sections: each a heading with bullets, entries, or
#                both. An entry is a title, an optional `detail` shown at the right
#                (the dates, or a project's stack), and bullets. The header (name,
#                contact, links) comes from profile.yaml.
"""

# The section headings a resume parser expects (the skill's ATS rules). A master
# writes them in capitals (Kim_Julongbayan_Resume.docx) or in title case
# (Kim_Julongbayan_Resume_AI.docx); any other bold paragraph is an entry.
STANDARD_HEADINGS = {
    "summary", "professional summary", "profile", "skills", "technical skills", "core skills",
    "experience", "work experience", "professional experience", "projects", "selected projects",
    "education", "certifications", "open source", "open source contributions", "awards",
    "publications", "languages", "volunteering", "volunteer experience",
}


def starter_spec(master: Path, company: str) -> str:
    """A spec holding the master resume's content, for the model to tailor.

    Reads the structure the masters in ~/port use: a bold heading (capitals, or a
    standard heading in any case), other bold paragraphs an entry ("title |
    detail"), List Bullet a bullet. The header is profile.yaml's, except the
    headline, which is the first line of the second paragraph.
    """
    import docx

    doc = docx.Document(str(master))
    paragraphs = doc.paragraphs
    headline = ""
    if len(paragraphs) > 1:
        headline = paragraphs[1].text.strip().splitlines()[0].strip() if paragraphs[1].text.strip() else ""
    sections: list[dict] = []
    entry = None
    for p in paragraphs[2:]:
        text = " ".join(p.text.split())
        if not text:
            continue
        bold = any(r.bold for r in p.runs if r.text.strip())
        if p.style.name.startswith("List"):
            if not sections:
                continue
            target = entry if entry is not None else sections[-1]
            target.setdefault("bullets", []).append(text)
        elif bold and " | " not in text and (text.isupper() or text.lower() in STANDARD_HEADINGS):
            sections.append({"heading": text})
            entry = None
        elif bold and sections:
            title, _, detail = text.rpartition(" | ") if " | " in text else (text, "", "")
            entry = {"title": title.strip(), "detail": detail.strip()} if detail else {"title": title.strip()}
            sections[-1].setdefault("entries", []).append(entry)
    spec = {
        "company": company,
        "company_names": [],
        "employer": "",
        "board": "none",
        "letter": "",
        "resume": {"headline": headline, "sections": sections},
    }
    body = yaml.safe_dump(spec, sort_keys=False, allow_unicode=True, width=10_000)
    return SPEC_HELP.format(company=company) + f"# Started from {master}\n" + body


# ── Building the documents ───────────────────────────────────────────────────


def _bare(url: str) -> str:
    return re.sub(r"^https?://(www\.)?", "", str(url).strip()).rstrip("/")


def header_lines(profile: dict) -> list[str]:
    top = "  |  ".join(str(profile[k]).strip() for k in ("location", "email", "phone")
                       if str(profile.get(k) or "").strip())
    links = "  |  ".join(f"{label}: {_bare(profile[key])}"
                         for label, key in (("Portfolio", "portfolio"), ("GitHub", "github"),
                                            ("LinkedIn", "linkedin"))
                         if str(profile.get(key) or "").strip())
    return [line for line in (top, links) if line]


def _run(paragraph, text, size, bold=False):
    from docx.shared import Pt

    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    return run


def _space(paragraph, before=0.0, after=0.0, align=None):
    from docx.shared import Pt

    fmt = paragraph.paragraph_format
    fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(before), Pt(after), 1.0
    if align is not None:
        fmt.alignment = align


def _rule_below(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    border, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
        bottom.set(qn(key), value)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)


def resume_docx(spec: dict, profile: dict, path: Path, stretch: float = 1.0) -> list[tuple[str, str]]:
    """Write the resume. Returns (kind, text) per paragraph, in page order, so the
    rendered PDF's lines can be traced back to the bullet they belong to."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    s = {k: v * stretch for k, v in SPACING.items()}
    doc = docx.Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(MARGINS_IN["top"]), Inches(MARGINS_IN["bottom"])
    section.left_margin, section.right_margin = Inches(MARGINS_IN["left"]), Inches(MARGINS_IN["right"])
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(SIZES["body"])
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(0), 1.0

    order: list[tuple[str, str]] = []
    center = WD_ALIGN_PARAGRAPH.CENTER

    def para(kind, text, size, bold=False, before=0.0, after=0.0, align=None):
        p = doc.add_paragraph()
        _space(p, before, after, align)
        _run(p, text, size, bold)
        order.append((kind, text))
        return p

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        _space(p, 0, s["bullet_after"])
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.first_line_indent = Inches(-0.13)
        _run(p, text, SIZES["body"])
        order.append(("bullet", text))

    para("name", str(profile.get("name") or "").upper(), SIZES["name"], bold=True, after=1, align=center)
    para("headline", spec["resume"]["headline"], SIZES["headline"], after=s["header_after"], align=center)
    for line in header_lines(profile):
        para("contact", line, SIZES["contact"], after=1, align=center)

    for sec in spec["resume"]["sections"]:
        _rule_below(para("heading", str(sec["heading"]).upper(), SIZES["heading"], bold=True,
                         before=s["heading_before"], after=s["heading_after"]))
        for text in sec.get("bullets") or []:
            bullet(str(text))
        for entry in sec.get("entries") or []:
            p = doc.add_paragraph()
            _space(p, s["entry_before"], s["entry_after"])
            _run(p, str(entry["title"]), SIZES["entry"], bold=True)
            detail = str(entry.get("detail") or "").strip()
            if detail:
                _run(p, "   |   " + detail, SIZES["entry_detail"])
            order.append(("entry", str(entry["title"]) + (" | " + detail if detail else "")))
            for text in entry.get("bullets") or []:
                bullet(str(text))
    doc.save(str(path))
    return order


def letter_docx(text: str, profile: dict, path: Path) -> None:
    """A cover letter that stands alone as an attachment: name, one contact line,
    the letter at 11pt. Short letters are meant to sit in white space."""
    import docx
    from docx.shared import Inches, Pt

    doc = docx.Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.9)
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(SIZES["letter"])

    head = doc.add_paragraph()
    _run(head, str(profile.get("name") or ""), 14, bold=True)
    head.paragraph_format.space_after = Pt(2)
    contact = doc.add_paragraph()
    _run(contact, "  |  ".join(header_lines(profile)), 9.5)
    contact.paragraph_format.space_after = Pt(14)
    for block in re.split(r"\n\s*\n", text.strip()):
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        # A sign-off block keeps its line breaks; a paragraph is one flowing line.
        if len(lines) > 1 and all(len(line) < 60 for line in lines):
            for n, line in enumerate(lines):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0 if n < len(lines) - 1 else 10)
                _run(p, line, SIZES["letter"])
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.15
        _run(p, " ".join(lines), SIZES["letter"])
    doc.save(str(path))


# ── Word or LibreOffice, then reading the page back ──────────────────────────


def _ps_quote(path: Path) -> str:
    return "'" + str(path).replace("'", "''") + "'"


def to_pdf(pairs: list[tuple[Path, Path]], timeout: int = 240) -> None:
    """Convert every (docx, pdf) pair in ONE Word session, or with LibreOffice.

    One session because starting Word is most of the cost, and the fit converts
    nine versions of the page. Convert inside a local temp folder: Word writing
    straight into a OneDrive folder hung with no window on 2026-09-11. A Word
    that stops answering is killed rather than left holding the next run.
    """
    if os.name == "nt":
        steps = "; ".join(
            f"$d = $w.Documents.Open({_ps_quote(src)}, $false, $true); "
            f"$d.SaveAs([ref]{_ps_quote(dst)}, [ref]17); $d.Close($false)"
            for src, dst in pairs)
        script = ("$ErrorActionPreference = 'Stop'; $w = New-Object -ComObject Word.Application; "
                  "$w.Visible = $false; $w.DisplayAlerts = 0; "
                  f"try {{ {steps} }} finally {{ $w.Quit() }}")
        started = time.time()
        try:
            done = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                                  capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            _kill_word_since(started)
            raise RenderError(f"Word did not finish within {timeout}s and was closed. Try again; "
                              "if it keeps happening, open Word once by hand to clear any dialog.")
        if done.returncode == 0 and all(dst.exists() for _, dst in pairs):
            return
        word_error = (done.stderr or done.stdout).strip().splitlines()[-1:] or ["no output"]
        if not shutil.which("soffice"):
            raise RenderError(f"Word could not make the PDF: {word_error[0]}")
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        raise RenderError("Nothing here can make a PDF: it needs Microsoft Word (Windows) or "
                          "LibreOffice (`soffice` on the PATH).")
    for src, dst in pairs:
        subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", str(dst.parent),
                        str(src)], capture_output=True, timeout=timeout, check=True)
        made = dst.parent / (src.stem + ".pdf")
        if made != dst:
            made.replace(dst)


def _kill_word_since(started: float) -> None:
    """Close only a Word this run started, never one the owner has open."""
    if os.name != "nt":
        return
    script = ("Get-Process WINWORD -ErrorAction SilentlyContinue | Where-Object "
              f"{{ $_.StartTime -ge (Get-Date).AddSeconds(-{int(time.time() - started) + 5}) }} | "
              "Stop-Process -Force")
    subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                   capture_output=True, timeout=30)


def pdf_lines(path: Path) -> list[dict]:
    """Every text line on every page, top to bottom: page, y (points from the
    bottom edge) and text."""
    from pypdf import PdfReader

    out = []
    for number, page in enumerate(PdfReader(str(path)).pages, 1):
        pieces = []

        def visit(text, cm, tm, _font, _size):
            if text and text.strip():
                x = tm[4] * cm[0] + tm[5] * cm[2] + cm[4]
                y = tm[4] * cm[1] + tm[5] * cm[3] + cm[5]
                pieces.append((y, x, text))

        page.extract_text(visitor_text=visit)
        pieces.sort(key=lambda p: (-p[0], p[1]))
        rows: list[list] = []
        for y, x, text in pieces:
            if rows and abs(rows[-1][0] - y) <= 2:
                rows[-1][1].append((x, text))
            else:
                rows.append([y, [(x, text)]])
        for y, parts in rows:
            out.append({"page": number, "y": y,
                        "text": "".join(t for _, t in sorted(parts)).strip()})
    return out


def _norm(text: str) -> str:
    return re.sub(r"[\W_]+", "", unicodedata.normalize("NFKC", text)).lower()


def line_counts(order: list[tuple[str, str]], lines: list[dict]) -> list[int] | None:
    """How many rendered lines each paragraph took, or None when the PDF's text
    cannot be followed back to the paragraphs (the wrap check is then skipped,
    and says so, rather than guessing)."""
    counts, at = [], 0
    for _kind, text in order:
        target, got, used = _norm(text), "", 0
        if not target:
            counts.append(0)
            continue
        while at < len(lines) and len(got) < len(target):
            got += _norm(lines[at]["text"])
            at += 1
            used += 1
        if got != target:
            return None
        counts.append(used)
    return counts


def measure(pdf: Path) -> dict:
    lines = pdf_lines(pdf)
    first = [line for line in lines if line["page"] == 1]
    bottom_pt = MARGINS_IN["bottom"] * 72
    lowest = min((line["y"] for line in first if line["y"] > 1), default=bottom_pt)
    return {"pages": max((line["page"] for line in lines), default=1),
            "waste_in": round((lowest - bottom_pt) / 72, 2),
            "overflow_lines": sum(1 for line in lines if line["page"] > 1),
            "lines": lines}


# ── The checks ───────────────────────────────────────────────────────────────


def _word_in(phrase: str, text: str) -> bool:
    return bool(re.search(r"(?<![\w-])" + re.escape(phrase) + r"(?![\w-])", text, re.I))


def _camel_words(token: str) -> str:
    return " ".join(re.findall(r"[A-Z]+(?![a-z])|[A-Z][a-z0-9]*|[a-z0-9]+", token))


def previous_names(spec: dict, settings: dict, prefix: str) -> tuple[list[str], list[str]]:
    """Names of employers applied to before, which must not survive into this one.

    Sure ones come from earlier specs' company_names. Possible ones come from the
    output folder's file names (Kim_Julongbayan_Resume_SparkMarketing.pdf ->
    "Spark Marketing"), two words or more, since one-word tokens there are often
    ordinary words ("Pipeline", "Scout").
    """
    own = {str(n).lower() for n in spec.get("company_names") or []} | {spec["company"].lower()}
    sure: set[str] = set()
    for path in Path(settings["specs_dir"]).glob("*.yaml"):
        if path.stem == spec["company"]:
            continue
        try:
            other = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
        except yaml.YAMLError:
            continue
        sure |= {str(n) for n in other.get("company_names") or [] if str(n).lower() not in own}
    maybe: set[str] = set()
    out_dir = Path(settings["output_dir"])
    if out_dir.is_dir():
        pattern = re.compile(rf"^{re.escape(prefix)}_[A-Za-z]+_([A-Za-z0-9]+)\.")
        for path in out_dir.iterdir():
            match = pattern.match(path.name)
            if not match or match.group(1) == spec["company"]:
                continue
            words = _camel_words(match.group(1))
            if len(words.split()) >= 2 and words.lower() not in own:
                maybe.add(words)
    return sorted(sure), sorted(maybe - sure)


URL = re.compile(r"(?<![@\w.])(?:https?://)?(?:[a-z0-9-]+\.)+(?:dev|io|com|ph|app|ai|net|org|co)"
                 r"(?:/[^\s,;)\"']*)?", re.I)
# Sites that refuse automated requests whatever the link: a failure there says
# nothing about the link, so they are named for a human to open instead.
UNCHECKABLE = ("linkedin.com",)


def check_links(texts: list[str], fetch=None) -> tuple[list[str], list[str]]:
    import httpx

    fetch = fetch or (lambda url: httpx.get(url, follow_redirects=True, timeout=12,
                                             headers={"User-Agent": "Mozilla/5.0 jobsift"}))
    found = []
    for text in texts:
        for match in URL.finditer(text):
            url = match.group(0).rstrip(".")
            if url.lower() not in (u.lower() for u in found):
                found.append(url)
    fails, warns = [], []
    for url in found:
        full = url if url.lower().startswith("http") else "https://" + url
        if any(host in full.lower() for host in UNCHECKABLE):
            warns.append(f"{url}: that site blocks automated checks, open it yourself")
            continue
        try:
            status = fetch(full).status_code
        except Exception as exc:
            fails.append(f"{url}: unreachable ({type(exc).__name__})")
            continue
        if status >= 400:
            fails.append(f"{url}: HTTP {status}")
    return fails, warns


def letter_paragraphs(letter: str) -> list[str]:
    return [" ".join(p.split()) for p in re.split(r"\n\s*\n", letter.strip()) if p.strip()]


def check_documents(spec: dict, docs: dict[str, str], rules: list, previous: tuple, links=None,
                    wraps=None) -> tuple[list[str], list[str]]:
    """Every check that needs no model. `docs` maps a document name to its text."""
    from .career import check_text

    fails, warns = [], []
    allowed = {w.lower() for w in spec.get("allow_words") or []}
    for name, text in docs.items():
        british = sorted({m.group(0) for m in BRITISH.finditer(text) if m.group(0).lower() not in allowed})
        if british:
            fails.append(f"{name}: British spelling: {', '.join(british)} (allow_words if a proper noun)")
        if "—" in text:
            fails.append(f"{name}: {text.count(chr(0x2014))} em dash(es); use a colon, comma or full stop")
        for hit in check_text(text, rules):
            fails.append(f"{name}: {hit['rule']}: {hit['problem']}. {hit['why']}")
        sure, maybe = previous
        for other in sure:
            if _word_in(other, text):
                fails.append(f"{name}: names {other!r}, an employer from an earlier application")
        for other in maybe:
            if _word_in(other, text):
                warns.append(f"{name}: mentions {other!r}, which is also an earlier application's name")

    letter = str(spec.get("letter") or "").strip()
    if letter:
        paragraphs = letter_paragraphs(letter)
        body = paragraphs[1:] if paragraphs and len(paragraphs[0]) < 60 and paragraphs[0].endswith(",") else paragraphs
        words = sum(len(p.split()) for p in body)
        low, high = LETTER_WORDS
        if not low <= words <= high:
            warns.append(f"letter: {words} words, aim for {low} to {high}")
        names = [str(n) for n in spec.get("company_names") or []]
        if names and not any(_word_in(n, letter) for n in names):
            fails.append(f"letter: never names the employer ({' / '.join(names)})")
        openings: dict[str, int] = {}
        for p in body:
            key = " ".join(p.lower().split()[:2])
            openings[key] = openings.get(key, 0) + 1
        repeated = [k for k, n in openings.items() if n > 1]
        if repeated:
            warns.append(f"letter: paragraphs open the same way ({', '.join(repr(k) for k in repeated)})")
        if re.search(r"\bso here is\b", letter, re.I):
            warns.append("letter: the 'so here is' hinge, which reads as a formula")
    for key in ("email", "message"):
        subject = str((spec.get(key) or {}).get("subject") or "")
        if len(subject) > SUBJECT_MAX:
            fails.append(f"{key}.subject: {len(subject)} characters, the limit is {SUBJECT_MAX}")
    if links:
        fails += links[0]
        warns += links[1]
    for w in wraps or []:
        warns.append(w)
    return fails, warns


# ── Putting it together ──────────────────────────────────────────────────────


def _email_text(spec: dict, profile: dict) -> str:
    from .draft import _with_signoff

    email = spec.get("email") or {}
    head = [f"TO: {email['to']}"] if email.get("to") else []
    head.append(f"SUBJECT: {email['subject']}")
    return "\n".join(head) + "\n\n" + _with_signoff(str(spec["letter"]).strip(), profile) + "\n"


def _message_text(spec: dict) -> str:
    message = spec.get("message") or {}
    link = str(message.get("link") or "").strip()
    return (f"SUBJECT: {message['subject']}\n\n{str(spec['letter']).strip()}\n\n"
            + (f"Resume: {link}" if link else LINK_PLACEHOLDER) + "\n")


def render(spec_file, profile: dict, settings: dict, rules: list, out_dir=None,
           offline: bool = False, convert=None, fetch=None) -> dict:
    """Render one application and check it. `convert` and `fetch` are for the dry
    run, which has no Word and no network."""
    spec = load_spec(spec_file)
    convert = convert or to_pdf
    out = Path(out_dir or settings["output_dir"]).expanduser()
    out.mkdir(parents=True, exist_ok=True)
    prefix = file_prefix(profile)
    company = spec["company"]
    written: list[Path] = []

    with tempfile.TemporaryDirectory(prefix="jobsift-render-") as tmp:
        tmp = Path(tmp)
        orders, pairs = {}, []
        for stretch in STRETCHES:
            src = tmp / f"resume-{stretch}.docx"
            orders[stretch] = resume_docx(spec, profile, src, stretch)
            pairs.append((src, tmp / f"resume-{stretch}.pdf"))
        letter_pdf = None
        if spec["board"] == "form" and str(spec.get("letter") or "").strip():
            from .draft import _with_signoff

            src = tmp / "letter.docx"
            letter_docx(_with_signoff(str(spec["letter"]).strip(), profile), profile, src)
            letter_pdf = tmp / "letter.pdf"
            pairs.append((src, letter_pdf))
        convert(pairs)

        results = {s: measure(tmp / f"resume-{s}.pdf") for s in STRETCHES}
        fitting = [s for s in STRETCHES if results[s]["pages"] == 1]
        # The fullest page that still fits; when none fits, the tightest, so the
        # report can say how far over it is.
        chosen = max(fitting) if fitting else min(STRETCHES)
        page = results[chosen]
        final = out / f"{prefix}_Resume_{company}.pdf"
        shutil.copyfile(tmp / f"resume-{chosen}.pdf", final)
        written.append(final)
        resume_text = "\n".join(line["text"] for line in page["lines"])

        docs = {final.name: resume_text}
        if letter_pdf is not None:
            dest = out / f"{prefix}_CoverLetter_{company}.pdf"
            shutil.copyfile(letter_pdf, dest)
            written.append(dest)
            letter_page = measure(letter_pdf)
            docs[dest.name] = "\n".join(line["text"] for line in letter_page["lines"])
        else:
            letter_page = None

    fails, warns = [], []
    if page["pages"] > 1:
        fails.append(f"resume: {page['pages']} pages at the tightest spacing, about "
                     f"{page['overflow_lines']} line(s) over. Cut that much; the three-line "
                     "bullets below are the first place to look")
    elif page["waste_in"] > MAX_WASTE_IN:
        free = int(page["waste_in"] * 72 // (SIZES["body"] * 1.22))
        warns.append(f"resume: {page['waste_in']}in empty under the last line at the loosest "
                     f"spacing, room for about {free} more line(s)")
    if letter_page is not None and letter_page["pages"] > 1:
        fails.append(f"cover letter: {letter_page['pages']} pages; a letter fits on one")

    wraps = []
    counts = line_counts(orders[chosen], page["lines"])
    if counts is None:
        wraps.append("resume: could not follow the PDF's lines back to the bullets, so the "
                      "three-line check was skipped")
    else:
        for (kind, text), n in zip(orders[chosen], counts):
            if kind == "bullet" and n >= 3:
                wraps.append(f"resume: a bullet runs to {n} lines, keep it to two: {text[:70]!r}...")

    if spec["board"] == "email":
        dest = out / f"{prefix}_Email_{company}.txt"
        dest.write_text(_email_text(spec, profile), encoding="utf-8")
        written.append(dest)
        docs[dest.name] = dest.read_text(encoding="utf-8")
    elif spec["board"] == "onlinejobs":
        dest = out / f"{prefix}_Message_{company}.txt"
        dest.write_text(_message_text(spec), encoding="utf-8")
        written.append(dest)
        docs[dest.name] = dest.read_text(encoding="utf-8")
        if not str((spec.get("message") or {}).get("link") or "").strip():
            warns.append("message: no resume link yet. Publish the resume with "
                         f"--publish {company}, which fills it in")

    links = None if offline else check_links(list(docs.values()), fetch=fetch)
    more_fails, more_warns = check_documents(spec, docs, rules,
                                             previous_names(spec, settings, prefix),
                                             links=links, wraps=wraps)
    return {"files": [str(p) for p in written], "pages": page["pages"],
            "waste_in": page["waste_in"], "stretch": chosen, "links_checked": not offline,
            "fail": fails + more_fails, "warn": warns + more_warns}


def report_text(report: dict) -> str:
    lines = [f"Wrote {len(report['files'])} file(s):"] + [f"  {f}" for f in report["files"]]
    lines.append(f"Resume: {report['pages']} page(s), {report['waste_in']}in left under the last "
                 f"line (spacing x{report['stretch']}).")
    if not report["links_checked"]:
        lines.append("Links were not checked (--offline).")
    if report["fail"]:
        lines += ["", f"FIX ({len(report['fail'])}):"] + [f"  - {f}" for f in report["fail"]]
    if report["warn"]:
        lines += ["", f"LOOK AT ({len(report['warn'])}):"] + [f"  - {w}" for w in report["warn"]]
    if not report["fail"] and not report["warn"]:
        lines += ["", "Every check passed."]
    elif not report["fail"]:
        lines += ["", "Nothing must be fixed; the notes above are judgment calls."]
    return "\n".join(lines)
